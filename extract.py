import os
from pptx import Presentation

pptx_path = '/Users/abhisheksrivastav/Documents/DR. Ritu Agrawal/DR RITU AGRAWAL FOR WEB SITE.pptx'
output_dir = '/Users/abhisheksrivastav/Documents/DR. Ritu Agrawal/public/assets/ppt_images'
text_output_file = '/Users/abhisheksrivastav/Documents/DR. Ritu Agrawal/ppt_content.txt'

os.makedirs(output_dir, exist_ok=True)

prs = Presentation(pptx_path)

with open(text_output_file, 'w', encoding='utf-8') as f:
    for i, slide in enumerate(prs.slides):
        f.write(f'--- Slide {i+1} ---\n')
        # Extract text
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                f.write(shape.text + '\n')
        
        # Extract images
        def extract_images_from_shape(shape):
            if shape.shape_type == 13: # Picture
                image = shape.image
                image_bytes = image.blob
                image_ext = image.ext
                image_filename = f'slide_{i+1}_{shape.shape_id}.{image_ext}'
                with open(os.path.join(output_dir, image_filename), 'wb') as img_file:
                    img_file.write(image_bytes)
            elif shape.shape_type == 6: # Group
                for child in shape.shapes:
                    extract_images_from_shape(child)
                    
        for shape in slide.shapes:
            extract_images_from_shape(shape)
        f.write('\n')

print("Extraction complete.")
